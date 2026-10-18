# ESTADOS DE CUENTA IBK

from __future__ import annotations

import re
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


# =========================
# RUTAS DE ENTRADA Y SALIDA
# =========================
BASE_DIR = Path(__file__).resolve().parent.parent
ENTRADA_DATOS = BASE_DIR / "Flujo" / "Estados_de_cuenta" / "EECC_IBK"
SALIDA_DIR = BASE_DIR / "Flujo" / "Output"
SALIDA_DATOS = SALIDA_DIR / "Movimientos_EECC_Interbank.xlsx"

# Expresiones para detectar las fechas, importes y número de cuenta.
PATRON_FECHA = re.compile(r"^\d{2}/\d{2}/\d{4}$")
PATRON_IMPORTE = re.compile(r"^[+-]?(?:\d{1,3}(?:,\d{3})*|\d+)\.\d{2}$")
PATRON_CUENTA = re.compile(
    r"(?:AHORRO|CUENTA)\s+SUELDO\s+SOLES\s*\n\s*(\d{3}-[\d-]{6,})",
    re.IGNORECASE,
)
PATRON_SALDO_INICIAL = re.compile(
    r"EMPEZASTE\s+\w+\s+CON\s+([+-]?[\d,]+\.\d{2})", re.IGNORECASE
)
PATRON_TOTALES = re.compile(
    r"SALDO\s+CONTABLE\s+AL\s+\d{2}/\d{2}\s+"
    r"([+-]?[\d,]+\.\d{2})\s+"
    r"([+-]?[\d,]+\.\d{2})\s+"
    r"([+-]?[\d,]+\.\d{2})",
    re.IGNORECASE,
)
CENTIMO = Decimal("0.01")

#  Transforma +1,200.50 o -20.00 a Decimal (sin errores binarios).
def monto_decimal(texto: str) -> Decimal:
    return Decimal(texto.replace(",", "")).quantize(CENTIMO)

# Excluye la última hoja informativa.
def es_pagina_de_ayuda(texto: str) -> bool:
    inicio = texto[:450].lower()
    return "te ayudamos a conocer" in inicio or "este es el ciclo de consumo" in inicio

# Agrupa palabras que comparten su coordenada vertical en el PDF.
def agrupar_palabras_por_renglon(palabras: list[dict], tolerancia: float = 1.6):
    grupos = []
    for palabra in sorted(palabras, key=lambda p: (p["top"], p["x0"])):
        if not grupos or abs(palabra["top"] - grupos[-1][0]["top"]) > tolerancia:
            grupos.append([palabra])
        else:
            grupos[-1].append(palabra)
    return [sorted(grupo, key=lambda p: p["x0"]) for grupo in grupos]

# Identifica cada movimiento por sus columnas (no por palabras de concepto).
def leer_movimientos_pagina(pagina, numero_pagina: int) -> list[dict]:
    palabras = pagina.extract_words()

    def cabecera(nombre: str):
        return next((p for p in palabras if p["text"].lower() == nombre.lower()), None)

    c_fecha = cabecera("Fecha")
    c_ingresos = cabecera("Ingresos")
    c_gastos = cabecera("Gastos")
    c_saldo = cabecera("Saldo")
    if not all((c_fecha, c_ingresos, c_gastos, c_saldo)):
        return []  # Primera página sin tabla o anexos publicitarios.

    x_ingresos = c_ingresos["x0"]
    x_gastos = c_gastos["x0"]
    x_saldo = c_saldo["x0"]
    limite_concepto = x_ingresos - 18
    limite_saldo = (x_gastos + x_saldo) / 2
    limite_ingreso_gasto = (x_ingresos + x_gastos) / 2

    movimientos = []
    for renglon in agrupar_palabras_por_renglon(palabras):
        # Una operación comienza en la primera columna con DD/MM/AAAA.
        fechas = [p for p in renglon if PATRON_FECHA.fullmatch(p["text"])
                  and p["x0"] < c_fecha["x0"] + 80
                  and p["top"] > c_fecha["top"] + 10]
        if not fechas:
            continue

        fecha_pdf = fechas[0]
        conceptos = [p["text"] for p in renglon
                     if p["x0"] > fecha_pdf["x1"] + 1
                     and p["x0"] < limite_concepto]

        importes = [p for p in renglon
                    if PATRON_IMPORTE.fullmatch(p["text"])
                    and p["x0"] >= limite_concepto]
        importes.sort(key=lambda p: p["x0"])

        if len(importes) != 2 or not conceptos:
            raise ValueError(
                f"Página {numero_pagina}, fecha {fecha_pdf['text']}: "
                "no se distinguieron concepto, movimiento y saldo. "
                "Revisar el formato del PDF."
            )
        casilla_movimiento, casilla_saldo = importes
        if not (casilla_movimiento["x0"] < limite_saldo
                <= casilla_saldo["x0"]):
            raise ValueError(
                f"Página {numero_pagina}: posiciones inesperadas de las cifras "
                f"en la fecha {fecha_pdf['text']}."
            )

        # El lugar de la cifra determina la columna, incluso si un PDF omite +/-.
        importe = abs(monto_decimal(casilla_movimiento["text"]))
        if casilla_movimiento["x0"] < limite_ingreso_gasto:
            movimiento = importe   # Ingresos
            tipo = "Ingreso"
        else:
            movimiento = -importe  # Gastos
            tipo = "Gasto"

        # Si el propio valor tiene signo, comprobar que no se contradiga con su columna.
        cifra_original = casilla_movimiento["text"]
        if ((cifra_original.startswith("-") and movimiento > 0)
                or (cifra_original.startswith("+") and movimiento < 0)):
            raise ValueError(
                f"Página {numero_pagina}, {fecha_pdf['text']}: "
                f"signo {cifra_original} contradice la columna {tipo}."
            )

        movimientos.append({
            "Fecha": datetime.strptime(fecha_pdf["text"], "%d/%m/%Y"),
            "Descripción": " ".join(conceptos),
            "Movimiento": movimiento,
            "Saldo del PDF": monto_decimal(casilla_saldo["text"]),
            "Página": numero_pagina,
        })
    return movimientos

# Comprueba saldos sucesivos y, si existen, totales impresos en el estado.
def validar_movimientos(registros: list[dict], texto_documento: str, archivo: Path):
    if not registros:
        raise ValueError("🚫 No se encontraron movimientos bancarios.")

    ingresos = sum((r["Movimiento"] for r in registros if r["Movimiento"] > 0), Decimal("0"))
    gastos = sum((r["Movimiento"] for r in registros if r["Movimiento"] < 0), Decimal("0"))

    saldo_inicial_match = PATRON_SALDO_INICIAL.search(texto_documento)
    if saldo_inicial_match:
        saldo_anterior = monto_decimal(saldo_inicial_match.group(1))
    else:
        # Permite validar saldos incluso si el documento no expresa el inicial.
        saldo_anterior = registros[0]["Saldo del PDF"] - registros[0]["Movimiento"]
        print(f"  Aviso: saldo inicial deducido a partir del primer movimiento ({saldo_anterior}).")

    saldo_calculado = saldo_anterior
    for fila, registro in enumerate(registros, start=1):
        saldo_calculado += registro["Movimiento"]
        if saldo_calculado.quantize(CENTIMO) != registro["Saldo del PDF"]:
            raise ValueError(
                f"Descuadre en {archivo.name}: movimiento {fila}, "
                f"{registro['Fecha']:%d/%m/%Y} - {registro['Descripción']}. "
                f"Saldo calculado {saldo_calculado:.2f} != "
                f"saldo PDF {registro['Saldo del PDF']:.2f}."
            )

    cierre_match = PATRON_TOTALES.search(texto_documento)
    if cierre_match:
        ingresos_pdf, gastos_pdf, cierre_pdf = map(monto_decimal, cierre_match.groups())
        if (ingresos != ingresos_pdf or gastos != gastos_pdf
                or saldo_calculado != cierre_pdf):
            raise ValueError(
                f"Totales no coinciden en {archivo.name}. "
                f"Ingresos: {ingresos:.2f} (PDF {ingresos_pdf:.2f}); "
                f"Gastos: {gastos:.2f} (PDF {gastos_pdf:.2f}); "
                f"saldo: {saldo_calculado:.2f} (PDF {cierre_pdf:.2f})."
            )
        print("  ✅ VALIDACIÓN OK: ingresos, gastos y saldo final coinciden con el PDF.")
    else:
        print("  ⚠️ Aviso: no se encontró la fila de totales; se validó el saldo de cada operación.")

    print(f"  📝 Registros: {len(registros)} | Ingresos: S/ {ingresos:,.2f} | "
          f"Gastos: S/ {gastos:,.2f} | Saldo final: S/ {saldo_calculado:,.2f}")
    return ingresos, gastos

# Lee páginas, extrae cuenta, fechas y movimientos; excluye guía de ayuda.
def procesar_pdf(ruta_pdf: Path) -> list[dict]:
    with pdfplumber.open(ruta_pdf) as pdf:
        if not pdf.pages:
            raise ValueError("El PDF no contiene páginas.")
        primera_hoja = pdf.pages[0].extract_text() or ""
        cuenta_match = PATRON_CUENTA.search(primera_hoja)
        if not cuenta_match:
            raise ValueError(" No se encontró el número bajo AHORRO SUELDO SOLES.")
        numero_cuenta = cuenta_match.group(1)

        todos = []
        texto_real = []
        for numero_pagina, pagina in enumerate(pdf.pages, start=1):
            texto = pagina.extract_text() or ""
            if es_pagina_de_ayuda(texto):
                print(f"  📄 Página {numero_pagina}: información ilustrativa, excluida.")
                continue
            texto_real.append(texto)
            filas = leer_movimientos_pagina(pagina, numero_pagina)
            todos.extend(filas)

    for registro in todos:
        registro["Número de cuenta"] = numero_cuenta
    validar_movimientos(todos, "\n".join(texto_real), ruta_pdf)
    return todos

# Crea una base lista para análisis: cuatro columnas, sin subtotales.
def exportar_excel(registros: list[dict], ruta_excel: Path):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Movimientos"
    hoja.append(["Fecha", "Número de cuenta", "Descripción", "Movimiento"])

    for registro in registros:
        hoja.append([
            registro["Fecha"],
            registro["Número de cuenta"],
            registro["Descripción"],
            float(registro["Movimiento"]),
        ])

    for celda in hoja[1]:
        celda.font = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="123B58")
        celda.alignment = Alignment(horizontal="center", vertical="center")
    hoja.row_dimensions[1].height = 26
    for columna, ancho in {"A": 19, "B": 24, "C": 36, "D": 19}.items():
        hoja.column_dimensions[columna].width = ancho
    for fila in hoja.iter_rows(min_row=2):
        fila[0].number_format = "DD/MM/YYYY"
        fila[0].alignment = Alignment(horizontal="center")
        fila[1].number_format = "@"  # Número de cuenta como texto
        fila[1].alignment = Alignment(horizontal="center")
        fila[3].number_format = '#,##0.00;[Red]-#,##0.00'
        fila[3].alignment = Alignment(horizontal="right")

    hoja.freeze_panes = "A2"
    if registros:
        # La tabla activa filtros y deja una única base limpia para BI.
        tabla = Table(displayName="TablaMovimientosIBK", ref=f"A1:D{len(registros)+1}")
        tabla.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False,
        )
        hoja.add_table(tabla)

    ruta_excel.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta_excel)


def main():
    print(f"\n")
    print("=" * 86)
    print(f"ESTADOS DE CUENTA DE INTERBANK")
    print("=" * 86)
    if not ENTRADA_DATOS.is_dir():
        raise FileNotFoundError(
            f" No existe la carpeta: {ENTRADA_DATOS}\n"
            "  Créala y coloca ahí los estados de cuenta PDF de Interbank."
        )
    pdfs = sorted(p for p in ENTRADA_DATOS.iterdir() if p.is_file() and p.suffix.lower() == ".pdf")
    if not pdfs:
        raise FileNotFoundError(f"❌ No se encontraron PDFs en: {ENTRADA_DATOS}")

    registros = []
    for pdf in pdfs:
        print(f"\n🔎 Archivo: {pdf.name}")
        filas = procesar_pdf(pdf)
        registros.extend(filas)

    # Orden cronológico global; empates mantienen orden de extracción del PDF.
    registros.sort(key=lambda r: r["Fecha"])

    ingresos = sum((r["Movimiento"] for r in registros if r["Movimiento"] > 0), Decimal(0))
    gastos = sum((r["Movimiento"] for r in registros if r["Movimiento"] < 0), Decimal(0))
    print("-" * 86)
    print(f"TOTAL INGRESOS : S/ {ingresos:,.2f}")
    print(f"TOTAL GASTOS   : S/ {gastos:,.2f}")
    print(f"MOVIMIENTO NETO: S/ {ingresos + gastos:,.2f}")
    exportar_excel(registros, SALIDA_DATOS)
    print(f"\n ✅ Excel generado exitosamente: {SALIDA_DATOS}")


if __name__ == "__main__":
    main()
