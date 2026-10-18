# RETO 1

from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# =========================
# RUTAS DE ENTRADA
# =========================
BASE_DIR = Path(__file__).resolve().parent.parent
ENTRADA_DATOS = BASE_DIR / "Flujo" / "Input" / "ventas_taller.csv"

# RUTA DE SALIDA
SALIDA_DIR = BASE_DIR / "Flujo" / "Output"
SALIDA_EXCEL = SALIDA_DIR / "Reporte_GlobalMarket.xlsx"


# =========================
# 0. CARGAR Y VALIDAR DATOS
# =========================
def cargar_datos():
    """Lee el CSV, comprueba las columnas y conserva las ventas de 2023."""
    if not ENTRADA_DATOS.is_file():
        raise FileNotFoundError(f"No se encontró el CSV en: {ENTRADA_DATOS}")

    # sep=None detecta el separador ',' o ';' más común en archivos CSV.
    df = pd.read_csv(ENTRADA_DATOS, sep=None, engine="python", encoding="utf-8-sig")
    df.columns = df.columns.str.strip().str.lower()

    columnas = [
        "fecha", "producto", "categoria", "region", "canal",
        "cliente_id", "ventas", "costo", "margen"
    ]
    faltantes = [col for col in columnas if col not in df.columns]
    if faltantes:
        raise ValueError(f"Faltan columnas en el CSV: {faltantes}")

    for col in ["producto", "categoria", "region", "canal"]:
        df[col] = df[col].astype("string").str.strip().replace("", pd.NA)

    fechas_texto = df["fecha"].astype("string").str.strip()
    formato_iso = fechas_texto.str.fullmatch(r"\d{4}-\d{2}-\d{2}").fillna(False).all()
    if formato_iso:
        df["fecha"] = pd.to_datetime(fechas_texto, format="%Y-%m-%d", errors="coerce")
    else:
        df["fecha"] = pd.to_datetime(fechas_texto, dayfirst=True, errors="coerce")

    for col in ["ventas", "costo", "margen"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    requeridas = ["fecha", "producto", "categoria", "region", "canal", "ventas", "costo", "margen"]
    conteo_nulos = df[requeridas].isna().sum()
    if conteo_nulos.sum() > 0:
        raise ValueError(
            "Hay valores vacíos o inválidos que deben revisarse:\n"
            + conteo_nulos[conteo_nulos > 0].to_string()
        )

    print("=" * 76)
    print("GLOBALMARKET - ANÁLISIS DE VENTAS 2023")
    print("=" * 76)
    print(f"Registros originales: {len(df)}")
    print(f"Posibles filas duplicadas (sin eliminar): {df.duplicated().sum()}")

    df = df.loc[df["fecha"].dt.year == 2023].copy()
    if df.empty:
        raise ValueError("El CSV no contiene ventas del año 2023.")
    print(f"Registros para 2023: {len(df)}")

    # Es solo una advertencia: se respeta el margen que viene en el CSV.
    diferencia = (df["ventas"] - df["costo"] - df["margen"]).abs()
    if (diferencia > 0.05).any():
        print("AVISO: algunas filas no cumplen ventas - costo = margen.")
        print("Se usará el campo margen original; revisa la definición del negocio.")
    return df


# =========================
# 1. VENTAS POR CATEGORÍA
# =========================
def tarea_1_ventas_por_categoria(df):
    """Suma ventas, costo y margen por categoría. Ordena por ventas."""
    tabla = (
        df.groupby("categoria", as_index=False)
        .agg(
            ventas_totales=("ventas", "sum"),
            costo_total=("costo", "sum"),
            margen_total=("margen", "sum")
        )
        .sort_values("ventas_totales", ascending=False)
    )

    print("\n" + "=" * 76)
    print("TAREA 1: VENTAS POR CATEGORÍA (USD)")
    print("=" * 76)
    print(tabla.round(2).to_string(index=False))
    return tabla


# =========================
# 2. DESEMPEÑO POR REGIÓN
# =========================
def tarea_2_desempeno_por_region(df):
    """Calcula ventas, cantidad de transacciones y margen medio por región."""
    tabla = (
        df.groupby("region", as_index=False)
        .agg(
            ventas_totales=("ventas", "sum"),
            numero_transacciones=("ventas", "size"),
            margen_promedio=("margen", "mean")
        )
        .sort_values("ventas_totales", ascending=False)
    )

    print("\n" + "=" * 76)
    print("TAREA 2: DESEMPEÑO POR REGIÓN (USD)")
    print("=" * 76)
    print(tabla.round(2).to_string(index=False))
    return tabla


# =========================
# 3. TOP 5 PRODUCTOS
# =========================
def tarea_3_top_5_productos(df):
    """Obtiene los cinco productos con mayor margen total y sus ventas."""
    tabla = (
        df.groupby(["producto", "categoria"], as_index=False)
        .agg(
            margen_total=("margen", "sum"),
            ventas_totales=("ventas", "sum")
        )
        .sort_values(["margen_total", "ventas_totales"], ascending=[False, False])
        .head(5)
    )

    print("\n" + "=" * 76)
    print("TAREA 3: TOP 5 PRODUCTOS POR MARGEN TOTAL (USD)")
    print("=" * 76)
    print(tabla.round(2).to_string(index=False))
    return tabla


# =========================
# 4. ANÁLISIS POR CANAL
# =========================
def tarea_4_analisis_por_canal(df):
    """Calcula participación de ventas (%) y margen promedio por canal."""
    tabla = (
        df.groupby("canal", as_index=False)
        .agg(
            ventas_totales=("ventas", "sum"),
            margen_promedio=("margen", "mean")
        )
    )
    ventas_globales = tabla["ventas_totales"].sum()
    if ventas_globales == 0:
        tabla["porcentaje_ventas"] = float("nan")
    else:
        tabla["porcentaje_ventas"] = tabla["ventas_totales"] / ventas_globales * 100

    tabla = tabla[["canal", "ventas_totales", "porcentaje_ventas", "margen_promedio"]]
    tabla = tabla.sort_values("ventas_totales", ascending=False)

    print("\n" + "=" * 76)
    print("TAREA 4: ANÁLISIS POR CANAL DE VENTA")
    print("=" * 76)
    print(tabla.round(2).to_string(index=False))
    print("Nota: porcentaje_ventas está expresado de 0 a 100; margen_promedio en USD.")
    return tabla


# =========================
# 5. TENDENCIA MENSUAL
# =========================
def tarea_5_tendencia_mensual(df):
    """Resume las ventas por mes y el cambio % respecto al mes anterior."""
    meses = pd.period_range("2023-01", "2023-12", freq="M")
    tabla = (
        df.groupby(df["fecha"].dt.to_period("M"))["ventas"]
        .sum()
        .reindex(meses, fill_value=0)
        .rename_axis("mes")
        .reset_index(name="ventas_totales")
    )
    ventas_anteriores = tabla["ventas_totales"].shift(1)
    # Cuando el mes previo tiene 0 ventas, no se puede calcular el crecimiento.
    tabla["variacion_pct"] = (
        (tabla["ventas_totales"] / ventas_anteriores.replace(0, float("nan"))) - 1
    ) * 100
    tabla["mes"] = tabla["mes"].astype(str)

    print("\n" + "=" * 76)
    print("TAREA 5: TENDENCIA MENSUAL DE VENTAS")
    print("=" * 76)
    print(tabla.round(2).to_string(index=False))
    print("Nota: NaN significa que no hay comparación válida con el mes anterior.")
    return tabla


# =========================
# 6. RETO FINAL: TABLA DINÁMICA
# =========================
def tarea_6_tabla_dinamica(df):
    """Crea una tabla dinámica por región y categoría; destaca el margen medio mayor."""
    tabla = (
        pd.pivot_table(
            df,
            index=["region", "categoria"],
            values=["ventas", "margen"],
            aggfunc={"ventas": "sum", "margen": "mean"},
            observed=True
        )
        .reset_index()
        .rename(columns={"ventas": "ventas_totales", "margen": "margen_promedio"})
        .sort_values("margen_promedio", ascending=False)
    )
    mejor = tabla.iloc[0]

    print("\n" + "=" * 76)
    print("TAREA 6: TABLA DINÁMICA POR REGIÓN Y CATEGORÍA")
    print("=" * 76)
    print(tabla.round(2).to_string(index=False))
    print(
        "\nSEGMENTO DESTACADO (mayor margen promedio por transacción): "
        f"{mejor['region']} / {mejor['categoria']} | USD {mejor['margen_promedio']:,.2f}"
    )
    return tabla


# =========================
# 7. EXPORTAR RESULTADOS A EXCEL
# =========================
def exportar_excel(resultados):
    """Guarda las seis tablas en hojas separadas con formato legible."""
    SALIDA_DIR.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(SALIDA_EXCEL, engine="openpyxl") as writer:
        for hoja, tabla in resultados.items():
            tabla.to_excel(writer, sheet_name=hoja, index=False)

    wb = load_workbook(SALIDA_EXCEL)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for celda in ws[1]:
            celda.fill = PatternFill("solid", fgColor="17365D")
            celda.font = Font(color="FFFFFF", bold=True)
            celda.alignment = Alignment(horizontal="center")

        for columna in ws.columns:
            letra = get_column_letter(columna[0].column)
            ancho = min(max(max(len(str(c.value or "")) for c in columna) + 2, 15), 30)
            ws.column_dimensions[letra].width = ancho

        for col_num, encabezado in enumerate(ws[1], start=1):
            nombre_col = str(encabezado.value)
            for fila in range(2, ws.max_row + 1):
                celda = ws.cell(row=fila, column=col_num)
                if nombre_col in ["porcentaje_ventas", "variacion_pct"]:
                    celda.number_format = '0.00"%"'  # El dato ya viene de 0 a 100.
                elif nombre_col == "numero_transacciones":
                    celda.number_format = "0"
                elif nombre_col in ["ventas_totales", "costo_total", "margen_total", "margen_promedio"]:
                    celda.number_format = '#,##0.00'

    # La primera fila es el segmento de mayor margen promedio.
    hoja_final = wb["06_Tabla_dinamica"]
    for celda in hoja_final[2]:
        celda.fill = PatternFill("solid", fgColor="DDF1D8")
        celda.font = Font(bold=True, color="14532D")
    hoja_final.cell(
        row=hoja_final.max_row + 2,
        column=1,
        value="Fila destacada: segmento con mayor margen promedio por transacción."
    )

    wb.save(SALIDA_EXCEL)
    print("\n" + "=" * 76)
    print(f"Excel guardado correctamente en:\n{SALIDA_EXCEL}")


# =========================
# 8. EJECUTAR TODO EL TRABAJO
# =========================
def main():
    datos = cargar_datos()

    # Cada función imprime su tabla y devuelve un DataFrame.
    resultados = {
        "01_Ventas_categoria": tarea_1_ventas_por_categoria(datos),
        "02_Desempeno_region": tarea_2_desempeno_por_region(datos),
        "03_Top5_productos": tarea_3_top_5_productos(datos),
        "04_Canales": tarea_4_analisis_por_canal(datos),
        "05_Tendencia_mensual": tarea_5_tendencia_mensual(datos),
        "06_Tabla_dinamica": tarea_6_tabla_dinamica(datos)
    }

    exportar_excel(resultados)


if __name__ == "__main__":
    main()
