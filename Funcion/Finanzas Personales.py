import os
from datetime import datetime
import openpyxl
from openpyxl import Workbook

import tkinter as tk
from tkinter import messagebox, Listbox
import customtkinter as ctk


# ======================================================
# RUTAS
# ======================================================
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

archivo_entrada_clientes = os.path.join(BASE_DIR, "Flujo", "input", "RUC-Clientes.xlsx")
archivo_salida = os.path.join(BASE_DIR, "Flujo", "output", "Base_actualizada_2.xlsx")


# ======================================================
# CONFIG EXCEL
# ======================================================
SHEETS = ["Ingresos", "Prestamos", "Gastos", "Presupuesto"]

COLUMNAS = {
    "Ingresos":    ["timestamp", "fecha", "cliente", "ruc", "categoria", "subcategoria", "monto", "moneda", "descripcion"],
    "Gastos":      ["timestamp", "fecha", "cliente", "ruc", "categoria", "subcategoria", "monto", "moneda", "descripcion"],
    "Prestamos":   ["timestamp", "fecha", "cliente", "ruc", "monto", "tasa", "plazo", "descripcion"],
    "Presupuesto": ["timestamp", "mes", "cliente", "ruc", "categoria", "monto"]
}

MONEDAS = ["PEN", "USD"]


# ======================================================
# HELPERS EXCEL
# ======================================================
def asegurar_excel_salida():
    os.makedirs(os.path.dirname(archivo_salida), exist_ok=True)

    if os.path.exists(archivo_salida):
        return

    wb = Workbook()
    ws = wb.active
    ws.title = SHEETS[0]
    ws.append(COLUMNAS[SHEETS[0]])

    for s in SHEETS[1:]:
        w = wb.create_sheet(s)
        w.append(COLUMNAS[s])

    wb.save(archivo_salida)
    wb.close()


def append_row(sheet, data):
    wb = openpyxl.load_workbook(archivo_salida)
    ws = wb[sheet]
    ws.append([data.get(c, "") for c in COLUMNAS[sheet]])
    wb.save(archivo_salida)
    wb.close()


# ======================================================
# CLIENTES (NO SE ELIMINA)
# ======================================================
def norm(txt):
    return str(txt).strip().lower()


def cargar_clientes_desde_excel():
    """
    Lee:
    Col A = RUC
    Col B = Cliente
    """
    clientes = []
    cliente_a_ruc = {}

    wb = openpyxl.load_workbook(archivo_entrada_clientes, data_only=True)
    ws = wb.active

    for r in range(2, ws.max_row + 1):
        ruc = ws.cell(r, 1).value
        cli = ws.cell(r, 2).value

        if cli:
            cli = str(cli).strip()
            ruc = "" if ruc is None else str(ruc).strip()
            clientes.append(cli)
            cliente_a_ruc[cli] = ruc

    wb.close()

    clientes = sorted(set(clientes), key=str.lower)
    return clientes, cliente_a_ruc


# ======================================================
# UI
# ======================================================
ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")


class App(ctk.CTk):

    def __init__(self):
        super().__init__()

        self.title("Registro de Flujo")
        self.geometry("820x560")
        self.resizable(False, False)

        asegurar_excel_salida()

        # --- clientes (NO DESAPARECE) ---
        self.clientes, self.cliente_a_ruc = cargar_clientes_desde_excel()

        # variables
        self.tipo = tk.StringVar(value="Gastos")
        self.fecha = tk.StringVar(value=datetime.now().strftime("%Y-%m-%d"))
        self.mes = tk.StringVar(value=datetime.now().strftime("%Y-%m"))

        self.cliente = tk.StringVar()
        self.ruc = tk.StringVar()

        self.categoria = tk.StringVar()
        self.subcategoria = tk.StringVar()
        self.monto = tk.StringVar()
        self.moneda = tk.StringVar(value="PEN")
        self.descripcion = tk.StringVar()

        # dropdown interno (NO Toplevel)
        self.dropdown_frame = None
        self.dropdown_listbox = None

        self._build_ui()

    # --------------------------------------------------
    def _build_ui(self):
        top = ctk.CTkFrame(self)
        top.pack(fill="x", padx=20, pady=10)

        ctk.CTkSegmentedButton(
            top,
            values=SHEETS,
            variable=self.tipo,
            command=lambda _: self._render_form()
        ).pack(fill="x")

        self.form = ctk.CTkFrame(self)
        self.form.pack(fill="both", expand=True, padx=20, pady=10)

        ctk.CTkButton(self, text="Guardar", command=self._guardar).pack(pady=10)

        self._render_form()

    # --------------------------------------------------
    def _clear_form(self):
        for w in self.form.winfo_children():
            w.destroy()
        self._hide_dropdown()

    # --------------------------------------------------
    def _render_form(self):
        self._clear_form()

        # Fecha / Mes
        if self.tipo.get() == "Presupuesto":
            ctk.CTkLabel(self.form, text="Mes").grid(row=0, column=0, sticky="w")
            ctk.CTkEntry(self.form, textvariable=self.mes).grid(row=0, column=1, sticky="we")
        else:
            ctk.CTkLabel(self.form, text="Fecha").grid(row=0, column=0, sticky="w")
            ctk.CTkEntry(self.form, textvariable=self.fecha).grid(row=0, column=1, sticky="we")

        # Cliente (autocomplete corregido)
        ctk.CTkLabel(self.form, text="Cliente").grid(row=1, column=0, sticky="w")
        self.entry_cliente = ctk.CTkEntry(self.form, textvariable=self.cliente)
        self.entry_cliente.grid(row=1, column=1, sticky="we")
        self.entry_cliente.bind("<KeyRelease>", self._autocomplete_cliente)

        # RUC
        ctk.CTkLabel(self.form, text="RUC").grid(row=2, column=0, sticky="w")
        ctk.CTkEntry(self.form, textvariable=self.ruc, state="disabled").grid(row=2, column=1, sticky="we")

        # Categoria
        ctk.CTkLabel(self.form, text="Categoría").grid(row=3, column=0, sticky="w")
        ctk.CTkEntry(self.form, textvariable=self.categoria).grid(row=3, column=1, sticky="we")

        # Subcategoria
        ctk.CTkLabel(self.form, text="Subcategoría").grid(row=4, column=0, sticky="w")
        ctk.CTkEntry(self.form, textvariable=self.subcategoria).grid(row=4, column=1, sticky="we")

        # Monto
        ctk.CTkLabel(self.form, text="Monto").grid(row=5, column=0, sticky="w")
        ctk.CTkEntry(self.form, textvariable=self.monto).grid(row=5, column=1, sticky="we")

        # Moneda
        ctk.CTkComboBox(self.form, values=MONEDAS, variable=self.moneda)\
            .grid(row=6, column=1, sticky="we")

        # Descripción
        ctk.CTkLabel(self.form, text="Descripción").grid(row=7, column=0, sticky="w")
        ctk.CTkEntry(self.form, textvariable=self.descripcion).grid(row=7, column=1, sticky="we")

        self.form.columnconfigure(1, weight=1)

    # --------------------------------------------------
    # AUTOCOMPLETE (CORREGIDO, SELECCIONABLE)
    # --------------------------------------------------
    def _autocomplete_cliente(self, event=None):
        texto = norm(self.cliente.get())

        if not texto:
            self._hide_dropdown()
            self.ruc.set("")
            return

        matches = [c for c in self.clientes if norm(c).startswith(texto)]
        if not matches:
            self._hide_dropdown()
            return

        self._show_dropdown(matches)

    def _show_dropdown(self, matches):
        self._hide_dropdown()

        self.dropdown_frame = ctk.CTkFrame(self.form, fg_color="#0b1220", corner_radius=6)
        self.dropdown_frame.place(in_=self.entry_cliente, relx=0, rely=1, relwidth=1)

        self.dropdown_listbox = Listbox(
            self.dropdown_frame,
            background="#0b1220",
            foreground="white",
            borderwidth=0,
            highlightthickness=0,
            selectbackground="#1e40af"
        )
        self.dropdown_listbox.pack(fill="both", expand=True)

        for m in matches[:8]:
            self.dropdown_listbox.insert("end", m)

        self.dropdown_listbox.bind("<ButtonRelease-1>", self._select_cliente)
        self.dropdown_listbox.bind("<Return>", self._select_cliente)
        self.dropdown_listbox.focus_set()

    def _hide_dropdown(self):
        if self.dropdown_frame:
            self.dropdown_frame.destroy()
            self.dropdown_frame = None
            self.dropdown_listbox = None

    def _select_cliente(self, event=None):
        if not self.dropdown_listbox:
            return

        sel = self.dropdown_listbox.curselection()
        if not sel:
            return

        nombre = self.dropdown_listbox.get(sel[0])
        self.cliente.set(nombre)
        self.ruc.set(self.cliente_a_ruc.get(nombre, ""))
        self._hide_dropdown()

    # --------------------------------------------------
    def _guardar(self):
        try:
            monto = float(self.monto.get())
        except:
            messagebox.showerror("Error", "Monto inválido")
            return

        data = {
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "fecha": self.fecha.get(),
            "mes": self.mes.get(),
            "cliente": self.cliente.get(),
            "ruc": self.ruc.get(),
            "categoria": self.categoria.get(),
            "subcategoria": self.subcategoria.get(),
            "monto": monto,
            "moneda": self.moneda.get(),
            "descripcion": self.descripcion.get()
        }

        append_row(self.tipo.get(), data)
        messagebox.showinfo("OK", "Registro guardado")

        self.cliente.set("")
        self.ruc.set("")
        self.monto.set("")


# ======================================================
if __name__ == "__main__":
    App().mainloop()
